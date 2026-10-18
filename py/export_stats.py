"""
export_stats.py
Xuất báo cáo trang Thống kê ra file Excel (.xlsx) hoặc PDF (.pdf).

  - Excel : openpyxl  (3 sheet: Tổng quan / Khu vực / Chi tiết lượt xe)
  - PDF   : reportlab (A4 ngang, dùng font tiếng Việt DejaVu đi kèm dự án)

Cách dùng (ở app.py):
    payload = export_stats.make_payload(stats, area_stats, records, ...)
    buf, mimetype, ext = export_stats.build('xlsx', payload)

Cả 2 thư viện đều nằm trong requirements.txt; nếu chưa cài thì hàm build()
ném ExportError để app.py hiện flash() hướng dẫn chạy lại setup.
"""

import os
from datetime import datetime
from io import BytesIO
from xml.sax.saxutils import escape

from config import FONTS_DIR

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_ERROR = None
except ImportError as _e:                 # pragma: no cover
    OPENPYXL_ERROR = _e

try:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (LongTable, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)
    REPORTLAB_ERROR = None
except ImportError as _e:                 # pragma: no cover
    REPORTLAB_ERROR = _e


TITLE = 'BÁO CÁO THỐNG KÊ BÃI ĐỖ XE'
REGULAR_FONT_NAME = 'VietFont'
BOLD_FONT_NAME = 'VietFont-Bold'

# Màu khớp biến --navy / --line trong theme.css (định dạng ARGB cho openpyxl)
X_NAVY = 'FF14306B'
X_LINE = 'FFE5E9F2'
X_ALT = 'FFF5F7FC'
X_TEXT = 'FF1D2433'
X_MUTED = 'FF7A8499'
X_GREEN = 'FF188246'
X_RED = 'FFB93636'

XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
PDF_MIME = 'application/pdf'


class ExportError(Exception):
    """Lỗi khi xuất file — app.py bắt lại và hiện qua flash()."""


# ============================================================
# DỮ LIỆU DÙNG CHUNG
# ============================================================

def fmt_money(value) -> str:
    return f'{int(value or 0):,}'.replace(',', '.') + ' đ'


def fmt_number(value) -> str:
    return f'{int(value or 0):,}'.replace(',', '.')


def _fmt_time(value) -> str:
    if not value:
        return '—'
    return str(value).replace('T', ' ')


def _fmt_date(value) -> str:
    """ISO 'YYYY-MM-DD' -> 'DD/MM/YYYY' (nếu không parse được thì trả về nguyên bản)."""
    if not value:
        return ''
    try:
        return datetime.strptime(str(value)[:10], '%Y-%m-%d').strftime('%d/%m/%Y')
    except ValueError:
        return str(value)


def _fmt_period(date_from, date_to) -> str:
    """Khoảng ngày dạng 'DD/MM/YYYY → DD/MM/YYYY'; nếu cùng một ngày thì chỉ in một lần."""
    start, end = _fmt_date(date_from), _fmt_date(date_to)
    if not start:
        return end
    if not end or start == end:
        return start
    return f'{start} → {end}'


def make_payload(stats, area_stats, records, range_label, date_from, date_to, generated_by):
    """Gom số liệu trang Thống kê + lượt xe trong kỳ thành cấu trúc dùng chung cho Excel/PDF.

    stats      : kết quả database.get_statistics()
    area_stats : kết quả database.area_stats()      ({'A': {...}, 'B': ..., 'C': ...})
    records    : kết quả database.get_records_between()
    """
    summary = [
        (0, 'Tổng doanh thu (từ trước đến nay)', stats['total_revenue'], 'money'),
        (1, 'Phí gửi xe', stats['parking_revenue'], 'money'),
        (1, 'Tiền vé tháng', stats['month_revenue'], 'money'),
        (0, f'Lượt vào {range_label}', stats['period_entries'], 'number'),
        (0, f'Lượt ra {range_label}', stats['period_exits'], 'number'),
        (0, f'Tổng lượt xe {range_label}', stats['period_records'], 'number'),
        (0, f'Doanh thu {range_label}', stats['period_revenue'], 'money'),
        (1, 'Phí gửi xe', stats['period_parking_revenue'], 'money'),
        (1, 'Tiền vé tháng', stats['period_month_revenue'], 'money'),
    ]

    areas = []
    for name in sorted(area_stats):
        d = area_stats[name]
        areas.append([name, d.get('total', 0), d.get('occupied', 0),
                      d.get('available', 0), d.get('maintenance', 0), d.get('percent', 0)])
    if areas:
        total = sum(a[1] for a in areas)
        occupied = sum(a[2] for a in areas)
        areas.append(['Tổng', total, occupied, sum(a[3] for a in areas),
                      sum(a[4] for a in areas),
                      round(occupied / total * 100, 1) if total else 0])

    rows = []
    for i, r in enumerate(records, 1):
        rows.append([
            i,
            r['plate_number'] or '',
            r['vehicle_type_name'] or '—',
            r['slot_name'] or '—',
            'Đang gửi' if r['status'] == 'IN' else 'Đã lấy xe',
            _fmt_time(r['entry_time']),
            _fmt_time(r['exit_time']),
            r['fee'] if r['fee'] is not None else None,
            r['note'] or '',
        ])

    return {
        'title': TITLE,
        'range_label': range_label,
        'date_from': date_from,
        'date_to': date_to,
        'generated_at': datetime.now().strftime('%d/%m/%Y %H:%M'),
        'generated_by': generated_by or '—',
        'summary': summary,
        'areas': areas,
        'records': rows,
    }


def build(fmt: str, payload: dict):
    """Trả về (BytesIO, mimetype, extension) theo fmt ('xlsx' | 'pdf')."""
    if fmt == 'xlsx':
        return build_xlsx(payload), XLSX_MIME, 'xlsx'
    if fmt == 'pdf':
        return build_pdf(payload), PDF_MIME, 'pdf'
    raise ExportError(f'Định dạng xuất file không hỗ trợ: {fmt}')


# ============================================================
# FONT TIẾNG VIỆT CHO PDF
# Ưu tiên font DejaVu đi kèm dự án (assets/fonts) -> chạy được trên mọi máy,
# không cần cài/cấu hình gì. Chỉ dự phòng font hệ thống nếu thiếu file đi kèm.
# ============================================================

def _font_roots():
    roots = [FONTS_DIR]                       # font đi kèm dự án (luôn ưu tiên)
    if os.name == 'nt':
        roots.append(os.path.join(os.environ.get('WINDIR', r'C:\Windows'), 'Fonts'))
    else:
        roots += [
            '/System/Library/Fonts/Supplemental',   # macOS (Arial, Times New Roman...)
            '/System/Library/Fonts',
            '/Library/Fonts',
            os.path.expanduser('~/Library/Fonts'),
            '/usr/share/fonts/truetype/dejavu',
            '/usr/share/fonts/truetype/liberation',
            '/usr/share/fonts/truetype/noto',
            '/usr/share/fonts/truetype',
            '/usr/share/fonts',
            os.path.expanduser('~/.fonts'),
        ]
    return [r for r in roots if os.path.isdir(r)]


def _font_names(bold: bool):
    """Tên file font tiếng Việt theo thứ tự ưu tiên (DejaVu đi kèm đứng đầu)."""
    if bold:
        return ['DejaVuSans-Bold.ttf', 'LiberationSans-Bold.ttf', 'NotoSans-Bold.ttf',
                'Arial Bold.ttf', 'Arial-Bold.ttf', 'arialbd.ttf', 'Tahoma Bold.ttf',
                'Verdana Bold.ttf', 'timesbd.ttf', 'tahomabd.ttf', 'verdanab.ttf']
    return ['DejaVuSans.ttf', 'LiberationSans-Regular.ttf', 'NotoSans-Regular.ttf',
            'Arial.ttf', 'arial.ttf', 'Times New Roman.ttf', 'times.ttf',
            'Tahoma.ttf', 'tahoma.ttf', 'Verdana.ttf', 'verdana.ttf']


def _find_font(bold: bool):
    """Trả về đường dẫn file .ttf tiếng Việt, hoặc None nếu không tìm thấy.
    Thứ tự: biến môi trường REPORT_FONT(_BOLD) -> font đi kèm dự án -> font hệ thống."""
    env_key = 'REPORT_FONT_BOLD' if bold else 'REPORT_FONT'
    env_path = (os.environ.get(env_key) or '').strip()
    if env_path and os.path.isfile(env_path):
        return env_path

    for root in _font_roots():
        for name in _font_names(bold):
            path = os.path.join(root, name)
            if os.path.isfile(path):
                return path

    # Không khớp tên nào -> quét mọi file .ttf có tên font phổ biến
    keywords = ('arial', 'dejavu', 'liberation', 'noto', 'tahoma', 'verdana', 'times')
    for root in _font_roots():
        for dirpath, _dirs, files in os.walk(root):
            for f in sorted(files):
                low = f.lower()
                if low.endswith('.ttf') and any(k in low for k in keywords):
                    if bold and not any(k in low for k in ('bold', '-bd', 'bd.')):
                        continue
                    return os.path.join(dirpath, f)
    return None


def _register_fonts():
    if REPORTLAB_ERROR is not None:
        raise ExportError('Chưa cài thư viện reportlab. Hãy chạy lại setup '
                          '(setup.bat / setup.sh) hoặc: pip install -r requirements.txt')
    if REGULAR_FONT_NAME in pdfmetrics.getRegisteredFontNames():
        return
    regular = _find_font(False)
    if not regular:
        raise ExportError(
            'Không tìm thấy font tiếng Việt để tạo PDF. File font đi kèm dự án bị thiếu '
            '(assets/fonts/DejaVuSans.ttf) và máy cũng không có font phù hợp. '
            'Hãy tải lại đầy đủ mã nguồn dự án, hoặc đặt biến môi trường '
            'REPORT_FONT=<đường dẫn file .ttf> rồi chạy lại ứng dụng.')
    bold = _find_font(True) or regular
    pdfmetrics.registerFont(TTFont(REGULAR_FONT_NAME, regular))
    pdfmetrics.registerFont(TTFont(BOLD_FONT_NAME, bold))
    pdfmetrics.registerFontFamily(REGULAR_FONT_NAME, normal=REGULAR_FONT_NAME,
                                  bold=BOLD_FONT_NAME, italic=REGULAR_FONT_NAME,
                                  boldItalic=BOLD_FONT_NAME)


# ============================================================
# EXCEL (.xlsx)
# ============================================================

def _border():
    side = Side(style='thin', color=X_LINE)
    return Border(left=side, right=side, top=side, bottom=side)


def _put(cell, value):
    """Gán giá trị cho ô Excel; chuỗi bắt đầu bằng = + - @ được giữ là văn bản
    để Excel không hiểu nhầm thành công thức."""
    cell.value = value
    if isinstance(value, str) and value[:1] in ('=', '+', '-', '@'):
        cell.data_type = 's'
    return cell


def _write_title(ws, payload, ncols):
    """3 dòng tiêu đề đầu báo cáo. Trả về hàng bắt đầu bảng (bỏ trống 1 dòng)."""
    end = get_column_letter(ncols)
    ws.merge_cells(f'A1:{end}1')
    ws['A1'] = payload['title']
    ws['A1'].font = Font(size=15, bold=True, color=X_NAVY)
    ws['A1'].alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 26

    ws.merge_cells(f'A2:{end}2')
    ws['A2'] = (f"Kỳ báo cáo: {payload['range_label']} "
                f"({_fmt_period(payload['date_from'], payload['date_to'])})")
    ws['A2'].font = Font(size=10.5, color=X_MUTED)

    ws.merge_cells(f'A3:{end}3')
    ws['A3'] = (f"Thời điểm xuất: {payload['generated_at']} "
                f"· Người xuất: {payload['generated_by']}")
    ws['A3'].font = Font(size=10.5, color=X_MUTED)
    return 5


def _write_header(ws, row, headers, widths=None):
    fill = PatternFill('solid', fgColor=X_NAVY)
    for col, text in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = Font(size=10.5, bold=True, color='FFFFFFFF')
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = _border()
    ws.row_dimensions[row].height = 24
    if widths:
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width


def _sheet_summary(wb, payload):
    ws = wb.active
    ws.title = 'Tổng quan'
    ws.sheet_view.showGridLines = False
    start = _write_title(ws, payload, 2)
    _write_header(ws, start, ['Chỉ tiêu', 'Giá trị'], widths=[46, 24])

    row = start + 1
    for indent, label, value, kind in payload['summary']:
        c1 = _put(ws.cell(row=row, column=1), label)
        c1.border = _border()
        c1.alignment = Alignment(vertical='center', indent=indent)
        if indent:
            c1.font = Font(size=11, italic=True, color=X_MUTED)

        c2 = _put(ws.cell(row=row, column=2), int(value or 0))
        c2.border = _border()
        c2.alignment = Alignment(horizontal='right', vertical='center')
        c2.number_format = '#,##0" đ"' if kind == 'money' else '#,##0'
        if indent:
            c2.font = Font(size=11, italic=True, color=X_MUTED)
        row += 1

    ws.freeze_panes = ws.cell(row=start + 1, column=1)


def _sheet_areas(wb, payload):
    ws = wb.create_sheet('Khu vực')
    ws.sheet_view.showGridLines = False
    start = _write_title(ws, payload, 6)
    headers = ['Khu vực', 'Tổng vị trí', 'Đang đỗ', 'Trống', 'Bảo trì', 'Tỉ lệ lấp đầy']
    _write_header(ws, start, headers, widths=[14, 14, 13, 11, 12, 18])

    row = start + 1
    for i, (area, total, occupied, available, maintenance, percent) in enumerate(payload['areas']):
        is_total = (str(area).lower() == 'tổng')
        values = [f'Khu {area}' if not is_total else 'Tổng cộng',
                  total, occupied, available, maintenance, percent]
        for col, value in enumerate(values, 1):
            cell = _put(ws.cell(row=row, column=col), value)
            cell.border = _border()
            if col == 1:
                cell.alignment = Alignment(horizontal='left', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='center', vertical='center')
                if col == 6:
                    cell.number_format = '0.0"%"'
                else:
                    cell.number_format = '#,##0'
            if is_total:
                cell.font = Font(bold=True, color=X_NAVY)
                cell.fill = PatternFill('solid', fgColor=X_ALT)
            elif i % 2:
                cell.fill = PatternFill('solid', fgColor=X_ALT)
        row += 1

    ws.freeze_panes = ws.cell(row=start + 1, column=1)


def _sheet_records(wb, payload):
    ws = wb.create_sheet('Chi tiết lượt xe')
    ws.sheet_view.showGridLines = False
    start = _write_title(ws, payload, 9)
    headers = ['STT', 'Biển số', 'Loại xe', 'Vị trí', 'Trạng thái',
               'Thời gian vào', 'Thời gian ra', 'Phí (VNĐ)', 'Ghi chú']
    _write_header(ws, start, headers,
                  widths=[6, 16, 16, 10, 13, 20, 20, 14, 42])

    row = start + 1
    if not payload['records']:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        cell = ws.cell(row=row, column=1, value='Không có dữ liệu trong kỳ.')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.font = Font(italic=True, color=X_MUTED)
        cell.border = _border()
    else:
        for i, rec in enumerate(payload['records']):
            for col, value in enumerate(rec, 1):
                cell = _put(ws.cell(row=row, column=col), value)
                cell.border = _border()
                cell.alignment = Alignment(vertical='center',
                                           horizontal='right' if col in (1, 8) else 'left',
                                           wrap_text=(col == 9))
                if col == 8 and value is not None:
                    cell.number_format = '#,##0" đ"'
                if col == 5:
                    cell.font = Font(bold=True,
                                     color=X_GREEN if value == 'Đang gửi' else X_RED)
                if i % 2:
                    cell.fill = PatternFill('solid', fgColor=X_ALT)
            row += 1
        ws.auto_filter.ref = f'A{start}:{get_column_letter(9)}{row - 1}'

    ws.freeze_panes = ws.cell(row=start + 1, column=1)


def build_xlsx(payload) -> BytesIO:
    if OPENPYXL_ERROR is not None:
        raise ExportError('Chưa cài thư viện openpyxl. Hãy chạy lại setup '
                          '(setup.bat / setup.sh) hoặc: pip install -r requirements.txt')
    wb = Workbook()
    _sheet_summary(wb, payload)
    _sheet_areas(wb, payload)
    _sheet_records(wb, payload)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ============================================================
# PDF (reportlab — A4 ngang)
# ============================================================

DETAIL_WIDTHS = [34, 74, 78, 40, 66, 110, 110, 62, 188]   # tổng 762pt (A4 ngang còn ~762pt)


def build_pdf(payload) -> BytesIO:
    _register_fonts()

    navy = colors.HexColor('#14306B')
    muted = colors.HexColor('#7A8499')
    line = colors.HexColor('#E5E9F2')
    alt = colors.HexColor('#F5F7FC')
    text = colors.HexColor('#1D2433')
    green = '#188246'
    red = '#B93636'

    st_title = ParagraphStyle('title', fontName=BOLD_FONT_NAME, fontSize=17, leading=21,
                              textColor=navy, spaceAfter=3)
    st_sub = ParagraphStyle('sub', fontName=REGULAR_FONT_NAME, fontSize=9.5, leading=13,
                            textColor=muted)
    st_h = ParagraphStyle('h', fontName=BOLD_FONT_NAME, fontSize=12, leading=15,
                          textColor=navy, spaceBefore=12, spaceAfter=6)
    st_cell = ParagraphStyle('cell', fontName=REGULAR_FONT_NAME, fontSize=8.5, leading=11.5,
                             textColor=text)
    st_cell_r = ParagraphStyle('cellr', parent=st_cell, alignment=2)
    st_cell_c = ParagraphStyle('cellc', parent=st_cell, alignment=TA_CENTER)
    st_small = ParagraphStyle('small', parent=st_cell, fontSize=7.5, leading=10)
    st_note = ParagraphStyle('note', parent=st_small, textColor=muted)
    st_th = ParagraphStyle('th', fontName=BOLD_FONT_NAME, fontSize=8.5, leading=11,
                           textColor=colors.white, alignment=TA_CENTER)

    # Bảng chi tiết có 9 cột -> dùng cỡ chữ nhỏ hơn cho vừa khổ A4 ngang
    st_d = ParagraphStyle('d', parent=st_cell, fontSize=8, leading=11)
    st_d_r = ParagraphStyle('dr', parent=st_d, alignment=2)
    st_d_c = ParagraphStyle('dc', parent=st_d, alignment=TA_CENTER)
    st_d_plate = ParagraphStyle('dp', parent=st_d_c, fontName=BOLD_FONT_NAME)
    st_d_note = ParagraphStyle('dn', parent=st_d, fontSize=7.5, leading=10, textColor=muted)
    st_th_d = ParagraphStyle('thd', parent=st_th, fontSize=8, leading=10.5)

    def table_style(nrows, first_data_row=1):
        style = [
            ('FONTNAME', (0, 0), (-1, -1), REGULAR_FONT_NAME),
            ('FONTSIZE', (0, 0), (-1, -1), 8.5),
            ('TEXTCOLOR', (0, 0), (-1, -1), text),
            ('BACKGROUND', (0, 0), (-1, 0), navy),
            ('FONTNAME', (0, 0), (-1, 0), BOLD_FONT_NAME),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.4, line),
            ('LINEBELOW', (0, 0), (-1, 0), 1, navy),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ]
        for i in range(first_data_row, nrows):
            if (i - first_data_row) % 2 == 1:
                style.append(('BACKGROUND', (0, i), (-1, i), alt))
        return TableStyle(style)

    page = landscape(A4)
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=page, leftMargin=14 * mm, rightMargin=14 * mm,
        topMargin=14 * mm, bottomMargin=16 * mm,
        title=payload['title'], author=payload['generated_by'])

    story = [
        Paragraph(payload['title'], st_title),
        Paragraph(f"Kỳ báo cáo: {payload['range_label']} "
                  f"({_fmt_period(payload['date_from'], payload['date_to'])})", st_sub),
        Paragraph(f"Thời điểm xuất: {payload['generated_at']} "
                  f"· Người xuất: {payload['generated_by']}", st_sub),
        Spacer(1, 6),
        Paragraph('1. TỔNG QUAN', st_h),
    ]

    # ----- Bảng tổng quan -----
    sum_data = [[Paragraph('Chỉ tiêu', st_th), Paragraph('Giá trị', st_th)]]
    for indent, label, value, kind in payload['summary']:
        shown = ('&nbsp;&nbsp;&nbsp;· ' if indent else '') + label
        number = fmt_money(value) if kind == 'money' else fmt_number(value)
        sum_data.append([Paragraph(shown, st_cell), Paragraph(number, st_cell_r)])
    sum_table = Table(sum_data, colWidths=[330, 170], repeatRows=1)
    sum_table.setStyle(table_style(len(sum_data)))
    story += [sum_table, Paragraph('2. TÌNH TRẠNG THEO KHU VỰC', st_h)]

    # ----- Bảng khu vực -----
    area_head = ['Khu vực', 'Tổng vị trí', 'Đang đỗ', 'Trống', 'Bảo trì', 'Tỉ lệ lấp đầy']
    area_data = [[Paragraph(h, st_th) for h in area_head]]
    for area, total, occupied, available, maintenance, percent in payload['areas']:
        is_total = str(area).lower() == 'tổng'
        name = 'Tổng cộng' if is_total else f'Khu {area}'
        row = [name, fmt_number(total), fmt_number(occupied), fmt_number(available),
               fmt_number(maintenance), f'{percent}%']
        area_data.append([Paragraph(f'<b>{v}</b>' if is_total else v, st_cell_c)
                          if i else Paragraph(f'<b>{v}</b>' if is_total else v, st_cell)
                          for i, v in enumerate(row)])
    area_table = Table(area_data, colWidths=[95, 90, 85, 85, 85, 110], repeatRows=1)
    area_style = table_style(len(area_data))
    if len(area_data) > 1:
        last = len(area_data) - 1
        area_style.add('BACKGROUND', (0, last), (-1, last), alt)
        area_style.add('LINEABOVE', (0, last), (-1, last), 0.9, navy)
    area_table.setStyle(area_style)
    story += [area_table, Paragraph(f'3. CHI TIẾT LƯỢT XE TRONG KỲ '
                                    f'({fmt_number(len(payload["records"]))} bản ghi)', st_h)]

    # ----- Bảng chi tiết lượt xe -----
    rec_head = ['STT', 'Biển số', 'Loại xe', 'Vị trí', 'Trạng thái',
                'Thời gian vào', 'Thời gian ra', 'Phí (VNĐ)', 'Ghi chú']
    rec_data = [[Paragraph(h, st_th_d) for h in rec_head]]
    for rec in payload['records']:
        stt, plate, vtype, slot, status, tin, tout, fee, note = rec
        status_color = green if status == 'Đang gửi' else red
        rec_data.append([
            Paragraph(str(stt), st_d_c),
            Paragraph(escape(plate), st_d_plate),
            Paragraph(escape(vtype), st_d),
            Paragraph(escape(slot), st_d_c),
            Paragraph(f'<font color="{status_color}">{escape(status)}</font>', st_d_c),
            Paragraph(escape(tin), st_d_c),
            Paragraph(escape(tout), st_d_c),
            Paragraph(fmt_money(fee) if fee is not None else '—', st_d_r),
            Paragraph(escape(note) if note and note != '—' else '—', st_d_note),
        ])
    if len(rec_data) == 1:
        rec_data.append([Paragraph('Không có dữ liệu trong kỳ.', st_d)] + [''] * 8)

    rec_table = LongTable(rec_data, colWidths=DETAIL_WIDTHS, repeatRows=1)
    rec_style = table_style(len(rec_data))
    if len(rec_data) == 2:
        rec_style.add('SPAN', (0, 1), (-1, 1))
        rec_style.add('ALIGN', (0, 1), (-1, 1), 'CENTER')
    rec_table.setStyle(rec_style)
    story.append(rec_table)

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont(REGULAR_FONT_NAME, 8)
        canvas.setFillColor(muted)
        canvas.drawString(doc_.leftMargin, 10 * mm,
                          f"{payload['title']} · Kỳ {payload['range_label']} · "
                          f"xuất ngày {payload['generated_at']}")
        canvas.drawRightString(page[0] - doc_.rightMargin, 10 * mm,
                               f'Trang {doc_.page}')
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    buf.seek(0)
    return buf
